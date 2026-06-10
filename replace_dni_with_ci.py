import os
from docx import Document

dirs = [
    r"c:\Users\karen\istpet-eventos",
    r"C:\Users\karen\.gemini\antigravity-ide\brain\b2ceec10-f8e4-4712-a939-2d2fd22c51b8"
]

def replace_in_paragraph(paragraph, old, new):
    # Case sensitive replace
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
        if old.lower() in run.text:
            run.text = run.text.replace(old.lower(), new.lower())
        if old.upper() in run.text:
            run.text = run.text.replace(old.upper(), new.upper())

def replace_in_table(table, old, new):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph, old, new)

def process_docx(filepath, old, new):
    try:
        doc = Document(filepath)
        for p in doc.paragraphs:
            replace_in_paragraph(p, old, new)
        for table in doc.tables:
            replace_in_table(table, old, new)
        doc.save(filepath)
        print(f"Modificado DOCX: {filepath}")
    except Exception as e:
        print(f"Error procesando {filepath}: {e}")

def process_md(filepath, old, new):
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        content = content.replace(old, new)
        content = content.replace(old.lower(), new.lower())
        content = content.replace(old.upper(), new.upper())
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)
        print(f"Modificado MD: {filepath}")
    except Exception as e:
        print(f"Error procesando {filepath}: {e}")

for d in dirs:
    for filename in os.listdir(d):
        filepath = os.path.join(d, filename)
        if filename.endswith(".docx") and not filename.startswith("~$"):
            process_docx(filepath, "DNI", "CI")
        elif filename.endswith(".md"):
            process_md(filepath, "DNI", "CI")
print("Reemplazo de DNI por CI completado en todos los manuales y guías.")
